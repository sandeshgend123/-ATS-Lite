"""Utility functions for application management"""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.core.files.base import ContentFile
from datetime import datetime


def generate_application_excel(application):
    """
    Generate Excel file with application details
    
    Args:
        application: Application instance
    
    Returns:
        ContentFile: Excel file as ContentFile object
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Application"
    
    # Define styles
    header_fill = PatternFill(start_color="0052CC", end_color="0052CC", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    section_fill = PatternFill(start_color="E7F0F7", end_color="E7F0F7", fill_type="solid")
    section_font = Font(bold=True, size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Set column widths
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 50
    
    row = 1
    
    # Title
    ws.merge_cells(f'A{row}:B{row}')
    title_cell = ws[f'A{row}']
    title_cell.value = "JOB APPLICATION DETAILS"
    title_cell.font = header_font
    title_cell.fill = header_fill
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    row += 2
    
    # Application ID Section
    ws[f'A{row}'].value = "Application ID"
    ws[f'B{row}'].value = application.id
    row += 1
    
    ws[f'A{row}'].value = "Submitted Date"
    ws[f'B{row}'].value = application.applied_at.strftime("%B %d, %Y at %I:%M %p")
    row += 1
    
    ws[f'A{row}'].value = "Last Updated"
    ws[f'B{row}'].value = application.updated_at.strftime("%B %d, %Y at %I:%M %p")
    row += 2
    
    # Candidate Information Section
    ws[f'A{row}'].value = "CANDIDATE INFORMATION"
    ws[f'A{row}'].font = section_font
    ws[f'A{row}'].fill = section_fill
    row += 1
    
    ws[f'A{row}'].value = "Name"
    ws[f'B{row}'].value = application.candidate.user.get_full_name() or application.candidate.user.username
    row += 1
    
    ws[f'A{row}'].value = "Email"
    ws[f'B{row}'].value = application.candidate.user.email
    row += 1
    
    ws[f'A{row}'].value = "Phone"
    ws[f'B{row}'].value = application.candidate.phone or "N/A"
    row += 1
    
    ws[f'A{row}'].value = "Bio"
    ws[f'B{row}'].value = application.candidate.bio or "N/A"
    ws[f'B{row}'].alignment = Alignment(wrap_text=True)
    row += 2
    
    # Job Information Section
    ws[f'A{row}'].value = "JOB INFORMATION"
    ws[f'A{row}'].font = section_font
    ws[f'A{row}'].fill = section_fill
    row += 1
    
    ws[f'A{row}'].value = "Job Title"
    ws[f'B{row}'].value = application.job.title
    row += 1
    
    ws[f'A{row}'].value = "Company"
    ws[f'B{row}'].value = application.job.company.user.get_full_name() or application.job.company.user.username
    row += 1
    
    ws[f'A{row}'].value = "Company Email"
    ws[f'B{row}'].value = application.job.company.user.email
    row += 1
    
    ws[f'A{row}'].value = "Job Description"
    ws[f'B{row}'].value = application.job.description
    ws[f'B{row}'].alignment = Alignment(wrap_text=True)
    row += 1
    
    ws[f'A{row}'].value = "Required Skills"
    ws[f'B{row}'].value = ", ".join(application.job.get_required_skills_list())
    ws[f'B{row}'].alignment = Alignment(wrap_text=True)
    row += 1
    
    ws[f'A{row}'].value = "Location"
    ws[f'B{row}'].value = application.job.location or "N/A"
    row += 1
    
    ws[f'A{row}'].value = "Salary Range"
    salary_range = "N/A"
    if application.job.salary_min and application.job.salary_max:
        salary_range = f"${application.job.salary_min:,} - ${application.job.salary_max:,}"
    elif application.job.salary_min:
        salary_range = f"${application.job.salary_min:,}+"
    ws[f'B{row}'].value = salary_range
    row += 1
    
    ws[f'A{row}'].value = "Experience Level"
    ws[f'B{row}'].value = application.job.experience_level.title()
    row += 1
    
    ws[f'A{row}'].value = "Job Type"
    ws[f'B{row}'].value = application.job.job_type.title()
    row += 2
    
    # Application Details Section
    ws[f'A{row}'].value = "APPLICATION DETAILS"
    ws[f'A{row}'].font = section_font
    ws[f'A{row}'].fill = section_fill
    row += 1
    
    ws[f'A{row}'].value = "Candidate Skills"
    ws[f'B{row}'].value = application.candidate_skills
    ws[f'B{row}'].alignment = Alignment(wrap_text=True)
    row += 1
    
    ws[f'A{row}'].value = "Skill Match Score"
    ws[f'B{row}'].value = f"{application.skill_match_score}%"
    row += 1
    
    ws[f'A{row}'].value = "Cover Letter"
    ws[f'B{row}'].value = application.cover_letter or "N/A"
    ws[f'B{row}'].alignment = Alignment(wrap_text=True)
    row += 1
    
    ws[f'A{row}'].value = "Resume File"
    ws[f'B{row}'].value = application.resume.name if application.resume else "N/A"
    row += 1
    
    ws[f'A{row}'].value = "Application Status"
    ws[f'B{row}'].value = application.get_status_display().title()
    row += 2
    
    # Footer
    ws.merge_cells(f'A{row}:B{row}')
    footer_cell = ws[f'A{row}']
    footer_cell.value = f"Generated on {datetime.now().strftime('%B %d, %Y at %I:%M %p')}"
    footer_cell.font = Font(italic=True, size=9, color="666666")
    footer_cell.alignment = Alignment(horizontal='center')
    
    # Generate file
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    
    # Create ContentFile
    filename = f"application_{application.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return ContentFile(excel_buffer.getvalue(), name=filename)
